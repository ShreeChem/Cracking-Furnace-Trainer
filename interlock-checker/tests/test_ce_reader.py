import pytest
from openpyxl import Workbook

from interlock_checker.ce_reader import CeReadError, read_ce


def test_synthetic_ce(seeded_files):
    ce = read_ce(seeded_files[1])
    assert len(ce.causes) == 12
    assert [e.tag for e in ce.effects][:3] == ["XV-100", "FV-100", "FV-110"]
    c = next(c for c in ce.causes if c.tag == "FZT-110")
    assert (c.voting, c.sil, c.latching, c.reset, c.setpoint) == ("2oo3", "SIL2", "YES", "AUTO", "0.20")
    xv172 = next(e for e in ce.effects if e.tag == "XV-172")
    assert (xv172.action, xv172.safe_position) == ("OPEN", "OPEN")
    # 'X' takes the column action; an explicit word is kept as written
    assert ce.links[(c.cause_id, "E01")] == "CLOSE"
    cot = next(c for c in ce.causes if c.tag == "TZT-150")
    assert ce.links[(cot.cause_id, "E04")] == "OPEN"
    assert ("C01", "E05") not in ce.links


def _write(tmp_path, rows, name="ce.xlsx"):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    path = tmp_path / name
    wb.save(path)
    return path


def test_alias_headings_and_sheet_choice(tmp_path):
    wb = Workbook()
    wb.active.title = "Cover"
    ws = wb.create_sheet("Matrix")
    for r in [["", "", "Final element", "XV-9"],
              ["", "", "Action", "close"],
              ["Cause No", "Tag", "Remarks", ""],
              ["1", "PZT-9", "", "x"],
              ["", "", "", ""],
              ["2", "TZT-9", "", "open"]]:
        ws.append(r)
    path = tmp_path / "alias.xlsx"
    wb.save(path)
    ce = read_ce(path)                          # finds the 'Matrix' sheet by itself
    assert [c.tag for c in ce.causes] == ["PZT-9", "TZT-9"]
    eff = ce.effects[0]
    assert eff.effect_id == "COL-D" and eff.action == "CLOSE"
    assert ce.links[("1", "COL-D")] == "CLOSE"
    assert ce.links[("2", "COL-D")] == "OPEN"
    assert read_ce(path, sheet="Matrix").causes
    with pytest.raises(CeReadError, match="Sheet 'Nope' not found"):
        read_ce(path, sheet="Nope")


def test_missing_file(tmp_path):
    with pytest.raises(CeReadError, match="not found"):
        read_ce(tmp_path / "x.xlsx")


def test_not_excel(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_text("hello")
    with pytest.raises(CeReadError, match="Could not open"):
        read_ce(p)


def test_no_header(tmp_path):
    with pytest.raises(CeReadError, match="Cause ID"):
        read_ce(_write(tmp_path, [["a", "b"]]))


def test_header_on_named_sheet_missing(tmp_path):
    with pytest.raises(CeReadError, match="no heading row"):
        read_ce(_write(tmp_path, [["a", "b"]]), sheet="Sheet")


def test_no_effect_label(tmp_path):
    with pytest.raises(CeReadError, match="Final element tag"):
        read_ce(_write(tmp_path, [["x"], ["Cause ID", "Initiator tag"], ["C1", "PZT-1"]]))


def test_no_effects(tmp_path):
    rows = [["", "Final element tag"], ["Cause ID", "Initiator tag"], ["C1", "PZT-1"]]
    with pytest.raises(CeReadError, match="No effect columns"):
        read_ce(_write(tmp_path, rows))


def test_no_causes(tmp_path):
    rows = [["", "Final element tag", "XV-1"], ["Cause ID", "Initiator tag", ""]]
    with pytest.raises(CeReadError, match="No cause rows"):
        read_ce(_write(tmp_path, rows))


def test_numbers_are_read_as_text(tmp_path):
    rows = [["", "", "Final element tag", "XV-1"], ["", "", "Action on trip", "CLOSE"],
            ["Cause ID", "Initiator tag", "Trip setpoint", ""], [1.0, "PZT-1", 62.0, "X"]]
    ce = read_ce(_write(tmp_path, rows))
    assert ce.causes[0].cause_id == "1" and ce.causes[0].setpoint == "62"
