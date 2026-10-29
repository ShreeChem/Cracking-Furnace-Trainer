"""End-to-end: the synthetic example through the CLI, and the committed files."""

import pytest
from openpyxl import load_workbook

from interlock_checker import cli
from interlock_checker.report import ReportContext, write_reports
from interlock_checker.synthetic import build_pid_xml, main as synth_main

from conftest import EXAMPLES, RULES

# Every seeded mismatch listed in the README, as (check number, tag).
EXPECTED = {
    ("1", "TZT-160"), ("1", "XV-176"),
    ("2", "FZT-101D"), ("2", "PT-172"),
    ("3", "FV-110"), ("3", "TV-150"), ("3", "XV-172"), ("3", "XV-100"),
    ("4", "PZT-141"), ("4", "FZT-110"), ("4", "BZS-171"),
}


def test_seeded_example_finds_every_planted_mismatch(seeded_files):
    findings, ctx = cli.run(*seeded_files, RULES)
    got = {(f.check.split(".")[0], f.tags[0]) for f in findings}
    assert got == EXPECTED
    assert [f.finding_id for f in findings] == [f"F-{i:03d}" for i in range(1, 12)]
    assert {f.check.split(".")[0] for f in findings} == {"1", "2", "3", "4"}
    assert ctx.banner.startswith("SYNTHETIC")
    assert ctx.placeholder_rules == ["R-PH-01", "R-PH-02", "R-PH-03"]


def test_clean_example_has_zero_findings(clean_files):
    findings, _ = cli.run(*clean_files, RULES)
    assert findings == []


def test_cli_writes_three_reports(seeded_files, tmp_path, capsys):
    code = cli.main(["--pid", str(seeded_files[0]), "--ce", str(seeded_files[1]),
                     "--rules", str(RULES), "--out", str(tmp_path)])
    assert code == 0
    out = capsys.readouterr().out
    assert "11 findings: 9 HIGH, 2 MEDIUM" in out and "placeholder" in out
    md = (tmp_path / "interlock_report.md").read_text(encoding="utf-8")
    html = (tmp_path / "interlock_report.html").read_text(encoding="utf-8")
    assert "SYNTHETIC – FOR TRAINING ONLY" in md and "F-011" in md
    assert "SYNTHETIC – FOR TRAINING ONLY" in html and "XV-172" in html
    wb = load_workbook(tmp_path / "interlock_report.xlsx")
    ws = wb["Findings"]
    assert ws["A1"].value.startswith("SYNTHETIC")
    assert [c.value for c in ws[5]] == ["Finding ID", "Check", "Tag(s)", "Severity",
                                        "Explanation", "Suggested action"]
    assert ws.max_row == 5 + 11
    assert wb["Summary"]["B2"].value == 11


def test_cli_fail_on(seeded_files, clean_files, tmp_path):
    args = ["--rules", str(RULES), "--out", str(tmp_path), "--fail-on", "HIGH"]
    assert cli.main(["--pid", str(seeded_files[0]), "--ce", str(seeded_files[1])] + args) == 1
    assert cli.main(["--pid", str(clean_files[0]), "--ce", str(clean_files[1])] + args) == 0


@pytest.mark.parametrize("which", ["pid", "ce", "rules"])
def test_cli_input_errors_exit_2(seeded_files, tmp_path, capsys, which):
    args = {"pid": str(seeded_files[0]), "ce": str(seeded_files[1]), "rules": str(RULES)}
    args[which] = str(tmp_path / "missing.file")
    code = cli.main([f"--{k}={v}" for k, v in args.items()] + ["--out", str(tmp_path)])
    assert code == 2
    assert "ERROR" in capsys.readouterr().err


def test_cli_without_rules_file(seeded_files, tmp_path):
    findings, ctx = cli.run(*seeded_files)
    assert not any(f.check.startswith("4.") for f in findings)
    assert ctx.placeholder_rules == []


def test_empty_report_and_warnings(tmp_path):
    ctx = ReportContext(pid_source="p", ce_source="c", warnings=["pyDEXPI error: lost a valve"])
    paths = write_reports([], ctx, tmp_path, stem="empty")
    md = paths[1].read_text(encoding="utf-8")
    assert "No mismatches found." in md and "lost a valve" in md
    assert "SYNTHETIC" not in md
    assert "No mismatches found." in paths[2].read_text(encoding="utf-8")


def test_committed_pid_examples_match_generator():
    assert (EXAMPLES / "synthetic_furnace_pid.xml").read_bytes() == build_pid_xml(True)
    assert (EXAMPLES / "synthetic_furnace_pid_clean.xml").read_bytes() == build_pid_xml(False)


def test_generator_cli(tmp_path, capsys):
    assert synth_main(["--out", str(tmp_path)]) == 0
    names = sorted(p.name for p in tmp_path.iterdir())
    assert names == ["synthetic_furnace_ce.xlsx", "synthetic_furnace_ce_clean.xlsx",
                     "synthetic_furnace_pid.xml", "synthetic_furnace_pid_clean.xml"]
    wb = load_workbook(tmp_path / "synthetic_furnace_ce.xlsx")
    for ws in wb.worksheets:                       # every sheet carries the stamp
        assert ws["A1"].value == "SYNTHETIC – FOR TRAINING ONLY"
        assert ws.oddHeader.center.text == "SYNTHETIC – FOR TRAINING ONLY"
