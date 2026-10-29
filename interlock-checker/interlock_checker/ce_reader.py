"""Step 2: read the cause-and-effect matrix from Excel.

The reader looks for the headings rather than fixed cell addresses, so a C&E
with extra title rows, different column widths or a few extra columns still
reads. What it needs:

* a heading row with at least 'Cause ID' and 'Initiator tag' (other cause
  columns are optional, see CAUSE_COLUMNS for the accepted names);
* above it, in one column, the effect labels 'Final element tag' and
  optionally 'Effect ID', 'Description', 'Action on trip', 'Required safe position';
* effect columns to the right of those labels;
* intersection cells holding either 'X' (use the column's action) or an
  action word such as CLOSE / OPEN / STOP.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import Cause, CeMatrix, Effect


class CeReadError(Exception):
    """Raised when the C&E workbook cannot be understood. The message says why."""


def _key(text) -> str:
    """Lower-case a heading and drop punctuation, so 'Trip Set-point' == 'trip setpoint'."""
    return "".join(ch for ch in str(text or "").lower() if ch.isalnum())


# Cause field -> accepted heading spellings (compared with _key()).
CAUSE_COLUMNS = {
    "cause_id": ["Cause ID", "Cause No", "Cause", "ID"],
    "tag": ["Initiator tag", "Initiator", "Tag", "Instrument tag", "Cause tag"],
    "description": ["Description", "Cause description", "Service"],
    "setpoint": ["Trip setpoint", "Setpoint", "Set point", "Trip value"],
    "units": ["Units", "Unit", "UoM", "Eng units"],
    "trip_type": ["Trip type", "Type", "Alarm type", "Trip"],
    "voting": ["Voting", "Vote", "MooN", "Voting logic"],
    "sil": ["SIL", "SIL rating", "Integrity level"],
    "latching": ["Latching", "Latched", "Latch"],
    "reset": ["Reset", "Reset type", "Reset mode"],
    "notes": ["Notes", "Remarks", "Comment", "Comments"],
}

# Effect field -> accepted label spellings in the label column.
EFFECT_LABELS = {
    "effect_id": ["Effect ID", "Effect No", "Effect"],
    "tag": ["Final element tag", "Final element", "Effect tag", "Tag"],
    "description": ["Description", "Effect description", "Service"],
    "action": ["Action on trip", "Action", "Trip action"],
    "safe_position": ["Required safe position", "Safe position", "Safe state"],
}

# Words in an intersection cell that mean "linked, use the column's default action".
_DEFAULT_MARKS = {"X", "●", "•", "Y", "YES", "1"}


def _lookup(table: dict[str, list[str]]) -> dict[str, str]:
    return {_key(alias): name for name, aliases in table.items() for alias in aliases}


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _find_sheet(wb, sheet: str | None):
    """Use the named sheet, else the first sheet that has a 'Cause ID' heading."""
    if sheet:
        if sheet not in wb.sheetnames:
            raise CeReadError(f"Sheet '{sheet}' not found. Sheets in the file: {wb.sheetnames}")
        return wb[sheet]
    for ws in wb.worksheets:
        if _find_header_row(ws) is not None:
            return ws
    raise CeReadError(
        "No sheet with a 'Cause ID' heading was found. The C&E needs a heading row with "
        "'Cause ID' and 'Initiator tag'."
    )


def _find_header_row(ws) -> int | None:
    cause_id_keys = {_key(a) for a in CAUSE_COLUMNS["cause_id"]}
    tag_keys = {_key(a) for a in CAUSE_COLUMNS["tag"]}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60)):
        keys = {_key(c.value) for c in row if c.value is not None}
        if keys & cause_id_keys and keys & tag_keys:
            return row[0].row
    return None


def read_ce(path: str | Path, sheet: str | None = None) -> CeMatrix:
    """Read an IEC 62881 style C&E matrix from an .xlsx file."""
    path = Path(path)
    if not path.is_file():
        raise CeReadError(f"C&E file not found: {path}")
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as exc:
        raise CeReadError(f"Could not open '{path.name}' as an Excel workbook: {exc}") from exc

    ws = _find_sheet(wb, sheet)
    header_row = _find_header_row(ws)
    if header_row is None:
        raise CeReadError(f"Sheet '{ws.title}' has no heading row with 'Cause ID' and 'Initiator tag'.")

    # Map cause columns from the heading row (first match wins).
    cause_lookup = _lookup(CAUSE_COLUMNS)
    cause_cols: dict[str, int] = {}
    for cell in ws[header_row]:
        name = cause_lookup.get(_key(cell.value))
        if name and name not in cause_cols:
            cause_cols[name] = cell.column

    # Find the effect label column: the cell above the heading row reading 'Final element tag'.
    effect_lookup = _lookup(EFFECT_LABELS)
    tag_keys = {_key(a) for a in EFFECT_LABELS["tag"]}
    label_col = None
    for row in ws.iter_rows(min_row=1, max_row=header_row - 1):
        for cell in row:
            if _key(cell.value) in tag_keys:
                label_col = cell.column
                break
        if label_col:
            break
    if label_col is None:
        raise CeReadError(
            f"No 'Final element tag' label found above row {header_row} on sheet '{ws.title}'. "
            "The effect columns need that label so the reader knows where they are."
        )

    label_rows: dict[str, int] = {}
    for r in range(1, header_row):
        name = effect_lookup.get(_key(ws.cell(row=r, column=label_col).value))
        if name and name not in label_rows:
            label_rows[name] = r

    ce = CeMatrix(source=str(path))

    # Effects: every column right of the label column with a final element tag.
    for col in range(label_col + 1, ws.max_column + 1):
        tag = _clean(ws.cell(row=label_rows["tag"], column=col).value)
        if not tag:
            continue
        values = {name: _clean(ws.cell(row=r, column=col).value) for name, r in label_rows.items()}
        letter = get_column_letter(col)
        ce.effects.append(Effect(
            effect_id=values.get("effect_id") or f"COL-{letter}",
            tag=tag,
            description=values.get("description", ""),
            action=values.get("action", "").upper(),
            safe_position=values.get("safe_position", "").upper(),
            column=letter,
        ))
    if not ce.effects:
        raise CeReadError(f"No effect columns found on sheet '{ws.title}'.")
    effect_cols = {e.column: e for e in ce.effects}

    # Causes: every row below the headings with an initiator tag.
    for r in range(header_row + 1, ws.max_row + 1):
        tag = _clean(ws.cell(row=r, column=cause_cols["tag"]).value)
        if not tag:
            continue
        values = {name: _clean(ws.cell(row=r, column=c).value) for name, c in cause_cols.items()}
        cause = Cause(
            cause_id=values.get("cause_id") or f"ROW-{r}",
            tag=tag,
            description=values.get("description", ""),
            setpoint=values.get("setpoint", ""),
            units=values.get("units", ""),
            trip_type=values.get("trip_type", "").upper(),
            voting=values.get("voting", ""),
            sil=values.get("sil", "").upper(),
            latching=values.get("latching", "").upper(),
            reset=values.get("reset", "").upper(),
            notes=values.get("notes", ""),
            row=r,
        )
        ce.causes.append(cause)
        for letter, eff in effect_cols.items():
            mark = _clean(ws[f"{letter}{r}"].value).upper()
            if not mark:
                continue
            ce.links[(cause.cause_id, eff.effect_id)] = eff.action if mark in _DEFAULT_MARKS else mark

    if not ce.causes:
        raise CeReadError(f"No cause rows (rows with an initiator tag) found on sheet '{ws.title}'.")
    return ce
