"""Step 7: write the mismatch report as Excel, Markdown and HTML.

The same list of findings goes into all three, so the Excel copy can be
filtered and assigned like an action tracker, the Markdown copy can sit in a
repository or an MOC record, and the HTML copy opens in any browser.
"""

from __future__ import annotations

import html
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import SEVERITIES, Finding

COLUMNS = ["Finding ID", "Check", "Tag(s)", "Severity", "Explanation", "Suggested action"]
_SEV_COLOURS = {"HIGH": "F4C7C3", "MEDIUM": "FCE8B2", "LOW": "D9EAD3", "INFO": "E7E6E2"}


@dataclass
class ReportContext:
    """Facts about the run printed at the top of every report."""

    pid_source: str
    ce_source: str
    rules_source: str = ""
    banner: str = ""                                   # e.g. SYNTHETIC – FOR TRAINING ONLY
    placeholder_rules: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    created: str = field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d %H:%M"))


def number_findings(findings: list[Finding]) -> list[Finding]:
    """Sort by check, then severity, then tag; give each a stable ID F-001, F-002, ..."""
    order = {s: i for i, s in enumerate(SEVERITIES)}
    ordered = sorted(findings, key=lambda f: (f.check.split(".")[0], order.get(f.severity, 9), f.check, f.tag_text))
    for i, f in enumerate(ordered, start=1):
        f.finding_id = f"F-{i:03d}"
    return ordered


def _rows(findings):
    return [[f.finding_id, f.check, f.tag_text, f.severity, f.explanation, f.suggested_action] for f in findings]


def _summary_lines(findings: list[Finding]) -> list[tuple[str, int]]:
    by_sev = Counter(f.severity for f in findings)
    return [(s, by_sev.get(s, 0)) for s in SEVERITIES]


def _by_check(findings: list[Finding]) -> list[tuple[str, int]]:
    c = Counter(f.check.split(" [")[0] for f in findings)
    return sorted(c.items())


def write_excel(findings: list[Finding], ctx: ReportContext, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    row = 1
    if ctx.banner:
        ws.cell(row=1, column=1, value=ctx.banner).font = Font(bold=True, color="FFFFFF", size=13)
        ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor="C8372D")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
        ws.oddHeader.center.text = ctx.banner
        ws.oddFooter.center.text = ctx.banner
        row = 2
    ws.cell(row=row, column=1, value="Interlock consistency report").font = Font(bold=True, size=13)
    ws.cell(row=row + 1, column=1, value=f"P&ID: {ctx.pid_source}   |   C&E: {ctx.ce_source}   |   "
                                         f"Rules: {ctx.rules_source or '(none)'}   |   Run: {ctx.created}")
    head = row + 3
    for k, name in enumerate(COLUMNS, start=1):
        c = ws.cell(row=head, column=k, value=name)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D0D7DE")
    for r, values in enumerate(_rows(findings), start=head + 1):
        for k, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=k, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=4).fill = PatternFill("solid", fgColor=_SEV_COLOURS.get(values[3], "FFFFFF"))
    for k, w in enumerate([10, 34, 16, 10, 90, 60], start=1):
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    if findings:
        ws.auto_filter.ref = f"A{head}:{get_column_letter(len(COLUMNS))}{head + len(findings)}"

    summary = wb.create_sheet("Summary")
    lines = [(ctx.banner, "")] if ctx.banner else []
    lines += [("Total findings", len(findings))] + _summary_lines(findings) + [("", "")]
    lines += [("By check", "")] + _by_check(findings)
    if ctx.placeholder_rules:
        lines += [("", ""), ("Placeholder rules still in use", ", ".join(ctx.placeholder_rules))]
    for w in ctx.warnings:
        lines.append(("Reader warning", w))
    for r, (a, b) in enumerate(lines, start=1):
        summary.cell(row=r, column=1, value=a).font = Font(bold=True)
        summary.cell(row=r, column=2, value=b)
    summary.column_dimensions["A"].width = 46
    summary.column_dimensions["B"].width = 80
    wb.save(path)
    return path


def _md_cell(text: str) -> str:
    return str(text).replace("|", "\\|").replace("\n", " ")


def write_markdown(findings: list[Finding], ctx: ReportContext, path: Path) -> Path:
    out = []
    if ctx.banner:
        out += [f"> **{ctx.banner}**", ""]
    out += ["# Interlock consistency report", "",
            f"- P&ID: `{ctx.pid_source}`", f"- C&E: `{ctx.ce_source}`",
            f"- Rules: `{ctx.rules_source or '(none)'}`", f"- Run: {ctx.created}", "",
            "## Summary", "", f"**{len(findings)} findings**: "
            + ", ".join(f"{n} {s}" for s, n in _summary_lines(findings)), "",
            "| Check | Findings |", "|---|---|"]
    out += [f"| {_md_cell(c)} | {n} |" for c, n in _by_check(findings)]
    if ctx.placeholder_rules:
        out += ["", f"**Placeholder rules still in use:** {', '.join(ctx.placeholder_rules)}. "
                    "Replace them with your site's latching and reset rules."]
    if ctx.warnings:
        out += ["", "**Reader warnings:**", ""] + [f"- {w}" for w in ctx.warnings]
    out += ["", "## Findings", "", "| " + " | ".join(COLUMNS) + " |", "|" + "---|" * len(COLUMNS)]
    out += ["| " + " | ".join(_md_cell(v) for v in row) + " |" for row in _rows(findings)]
    if not findings:
        out += ["", "No mismatches found."]
    path.write_text("\n".join(out) + "\n", encoding="utf-8")
    return path


def write_html(findings: list[Finding], ctx: ReportContext, path: Path) -> Path:
    e = html.escape
    sev_css = "".join(f".s-{s}{{background:#{c}}}" for s, c in _SEV_COLOURS.items())
    rows = "".join(
        "<tr>" + "".join(f'<td class="s-{e(v)}">{e(v)}</td>' if i == 3 else f"<td>{e(str(v))}</td>"
                         for i, v in enumerate(r)) + "</tr>"
        for r in _rows(findings))
    summary = ", ".join(f"{n} {s}" for s, n in _summary_lines(findings))
    checks = "".join(f"<li>{e(c)}: {n}</li>" for c, n in _by_check(findings))
    extra = ""
    if ctx.placeholder_rules:
        extra += (f"<p class='warn'><b>Placeholder rules still in use:</b> {e(', '.join(ctx.placeholder_rules))}. "
                  "Replace them with your site's latching and reset rules.</p>")
    if ctx.warnings:
        extra += "<p><b>Reader warnings:</b></p><ul>" + "".join(f"<li>{e(w)}</li>" for w in ctx.warnings) + "</ul>"
    banner = f"<div class='banner'>{e(ctx.banner)}</div>" if ctx.banner else ""
    doc = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Interlock consistency report</title>
<style>
body{{font-family:system-ui,sans-serif;margin:16px;color:#1d1d1b;background:#fff}}
.banner{{background:#C8372D;color:#fff;font-weight:700;padding:8px 12px;text-align:center}}
table{{border-collapse:collapse;width:100%;font-size:14px}}
th,td{{border:1px solid #ccc;padding:6px;vertical-align:top;text-align:left}}
th{{background:#D0D7DE;position:sticky;top:0}}
.warn{{background:#FCE8B2;padding:8px}} .wrap{{overflow-x:auto}}
td:nth-child(1),td:nth-child(3){{white-space:nowrap}}
{sev_css}
</style></head><body>
{banner}
<h1>Interlock consistency report</h1>
<p>P&amp;ID: <code>{e(ctx.pid_source)}</code><br>C&amp;E: <code>{e(ctx.ce_source)}</code><br>
Rules: <code>{e(ctx.rules_source or '(none)')}</code><br>Run: {e(ctx.created)}</p>
<h2>Summary</h2><p><b>{len(findings)} findings</b>: {e(summary)}</p><ul>{checks}</ul>{extra}
<h2>Findings</h2>
<div class="wrap"><table><thead><tr>{''.join(f'<th>{c}</th>' for c in COLUMNS)}</tr></thead>
<tbody>{rows or '<tr><td colspan="6">No mismatches found.</td></tr>'}</tbody></table></div>
{banner}
</body></html>
"""
    path.write_text(doc, encoding="utf-8")
    return path


def write_reports(findings: list[Finding], ctx: ReportContext, out_dir: str | Path,
                  stem: str = "interlock_report") -> list[Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    return [
        write_excel(findings, ctx, out / f"{stem}.xlsx"),
        write_markdown(findings, ctx, out / f"{stem}.md"),
        write_html(findings, ctx, out / f"{stem}.html"),
    ]
